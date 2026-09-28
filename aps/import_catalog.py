"""Source-traceable readers and conservative candidates for factory workbooks."""
from __future__ import annotations

from decimal import Decimal, ROUND_HALF_UP
from fractions import Fraction
from hashlib import sha256
from io import BytesIO
import json
import math
import re
import unicodedata

import pandas as pd


def text(value):
    return "" if value is None or pd.isna(value) else str(value).strip()


def number(value):
    try:
        result = float(text(value).lstrip("xX×"))
        return result if math.isfinite(result) and result > 0 else None
    except ValueError:
        return None


def rounded(value):
    # Remove spreadsheet floating point noise before decimal half-up rounding.
    return Decimal(str(round(float(value), 9))).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def machine_codes(value):
    value = re.sub(r"\s+", "", text(value).upper())
    result = []
    for match in re.finditer(r"([AC])(\d+(?:[,，、/]\d+)*)", value):
        result.extend(match[1] + digit for digit in re.split(r"[,，、/]", match[2]))
    return sorted(set(result))


def raw_sheets(content: bytes, filename: str):
    """Expand actual merged cells only, never arbitrary blank cells."""
    if filename.lower().endswith(".xls"):
        import xlrd
        book = xlrd.open_workbook(file_contents=content, formatting_info=True)
        result = {}
        for sheet in book.sheets():
            rows = [sheet.row_values(i) for i in range(sheet.nrows)]
            for r0, r1, c0, c1 in sheet.merged_cells:
                for r in range(r0, min(r1, sheet.nrows)):
                    for c in range(c0, min(c1, sheet.ncols)):
                        rows[r][c] = rows[r0][c0]
            result[sheet.name] = pd.DataFrame(rows).fillna("")
        return result
    from openpyxl import load_workbook
    book = load_workbook(BytesIO(content), data_only=True)
    result = {}
    for sheet in book:
        rows = [list(row) for row in sheet.values]
        for region in sheet.merged_cells.ranges:
            for r in range(region.min_row - 1, region.max_row):
                for c in range(region.min_col - 1, region.max_col):
                    rows[r][c] = rows[region.min_row - 1][region.min_col - 1]
        result[sheet.title] = pd.DataFrame(rows).fillna("")
    book.close()
    return result


def reference(kind, sheet, row):
    return f"{kind}|{sheet}|{row}"


def read_capacity(content, filename):
    records, warnings = [], []
    for sheet, frame in raw_sheets(content, filename).items():
        locations = [(i, j) for i, row in frame.head(12).iterrows()
                     for j, value in enumerate(row) if text(value) == "實際產速"]
        if not locations:
            warnings.append(f"{sheet}：未發現實際產速欄，未作為產速資料")
            continue
        if len(locations) != 1:
            warnings.append(f"{sheet}：實際產速標題不唯一，需確認格式")
            continue
        header, speed_col = locations[0]
        labels = {text(v): j for j, v in enumerate(frame.iloc[header]) if text(v)}
        sub = {text(v): j for j, v in enumerate(frame.iloc[header + 1]) if text(v)}
        if not {"規格", "品名"}.issubset(sub):
            warnings.append(f"{sheet}：缺少規格／品名標題，需確認格式")
            continue
        def cell(row, label, mapping=labels):
            return row.iloc[mapping[label]] if label in mapping else ""
        unit = text(frame.iloc[header + 1, speed_col]).replace(" ", "").upper()
        for idx, row in frame.iloc[header + 2:].iterrows():
            name, spec = text(cell(row, "品名", sub)), text(cell(row, "規格", sub))
            if not name and not spec:
                continue
            primary, backup = machine_codes(cell(row, "生產")), machine_codes(cell(row, "備用"))
            stage = " ".join(text(v) for v in row.iloc[:3])
            reason = ""
            if unit not in {"M/分", "M/MIN"}:
                reason = "實際產速單位無法確認為 M/分"
            elif re.search(r"[內外]\s*管", stage):
                reason = "內管／外管分段製程，需另行確認完整製程產速"
            elif len(primary) > 1:
                reason = "同列列有多台生產機台，需提供各機台實際產速"
            elif not primary:
                reason = "未記載生產機台"
            records.append(dict(ref=reference("產能", sheet, idx + 1), sheet=sheet, row=idx + 1,
                                name=name, spec=spec, inner=number(cell(row, "內徑")),
                                outer=number(cell(row, "外徑")), thickness=number(cell(row, "厚度")),
                                speed=number(row.iloc[speed_col]), machines=primary, backup=backup,
                                reason=reason))
    if not records:
        raise ValueError("機台產能資料庫未讀到任何產能資料；請檢查工作表標題與格式")
    return records, warnings


def read_molds(content, filename):
    records, warnings = [], []
    for sheet, frame in raw_sheets(content, filename).items():
        locations = [(i, j) for i, row in frame.head(12).iterrows()
                     for j, value in enumerate(row) if text(value) == "成品規格"]
        if not locations:
            warnings.append(f"{sheet}：非成品規格模具表，未作為機台限制資料")
            continue
        header, col = locations[0]
        machine_cols = [j for _, row in frame.head(header + 3).iterrows() for j, v in enumerate(row)
                        if re.sub(r"\s+", "", text(v)) in {"生產機台", "適用產線"}]
        for idx, row in frame.iloc[header + 3:].iterrows():
            inner = number(row.iloc[col])
            if inner is None or col + 1 >= len(row):
                continue
            next_value = text(row.iloc[col + 1])
            if next_value.lower() in {"x", "×", "*"} and col + 2 < len(row):
                thickness = number(row.iloc[col + 2])
            elif next_value.lower().startswith("x"):
                thickness = number(next_value)
            else:
                continue
            if thickness is None:
                continue
            machines = machine_codes(row.iloc[machine_cols[0]]) if machine_cols else machine_codes(sheet)
            records.append(dict(ref=reference("模具", sheet, idx + 1), sheet=sheet, row=idx + 1,
                                name=text(row.iloc[col - 1]) if col else "", inner=inner,
                                thickness=thickness, machines=machines,
                                detail=" | ".join(f"{j + 1}:{text(v)}" for j, v in enumerate(row) if text(v))))
    if not records:
        raise ValueError("模具資料庫未讀到可辨識的成品規格；請檢查工作表格式")
    return records, warnings


NUM = r"(?:\d+[ -]+\d+/\d+|\d+/\d+|\d+(?:\.\d+)?)"


def dimension(value, unit):
    parts = re.split(r"[ -]+", value.strip())
    result = sum(Fraction(part) for part in parts)
    return float(result * (Fraction("25.4") if unit == '"' else 1))


def _parse_spec(spec):
    spec = unicodedata.normalize("NFKC", spec).upper().replace("’", "'").replace("‘", "'").replace("″", '"')
    result = {}
    for label, key in [("ID", "inner"), ("OD", "outer")]:
        match = re.search(label + r"\s*(" + NUM + r")\s*(MM|\")", spec)
        if not match:
            match = re.search(r"(" + NUM + r")\s*(MM|\")\s*" + label, spec)
        if match:
            result[key] = dimension(match[1], match[2])
    thick = re.search(r"厚\s*(" + NUM + r")\s*MM", spec)
    if thick:
        result["thickness"] = dimension(thick[1], "MM")
    if "inner" in result and "outer" in result:
        result["thickness"] = (result["outer"] - result["inner"]) / 2
    length = re.search(r"(?:X|×|\*)\s*(\d+(?:\.\d+)?)\s*(FT|'|M)\s*$", spec)
    if length:
        result["length_m"] = float(length[1]) * (0.3048 if length[2] in {"FT", "'"} else 1)
    if "outer" not in result and {"inner", "thickness"}.issubset(result):
        result["outer"] = result["inner"] + 2 * result["thickness"]
    return result


def parse_spec(spec):
    try:
        result = _parse_spec(spec)
        return {key: value for key, value in result.items() if math.isfinite(value) and value > 0}
    except (ValueError, ZeroDivisionError, OverflowError):
        return {}


def matching_dimensions(spec, record, mold=False):
    fields = ["inner", "thickness"] if mold else ["inner", "outer"]
    return all(spec.get(key) and record.get(key) and rounded(spec[key]) == rounded(record[key]) for key in fields)


def mapping_key(row):
    return sha256(json.dumps([text(row[k]) for k in ["產品品號", "品名", "規格"]], ensure_ascii=False).encode()).hexdigest()


def candidate_records(row, records, mold=False):
    from difflib import SequenceMatcher
    spec = parse_spec(row["規格"])
    candidates = [r for r in records if matching_dimensions(spec, r, mold)]
    # Similarity sorts suggestions only. It never authorizes a match.
    return sorted(candidates, key=lambda r: SequenceMatcher(None, text(row["品名"]), r["name"]).ratio(), reverse=True)
