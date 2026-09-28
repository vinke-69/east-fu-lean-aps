"""Read exported orders without depending on visible Excel columns."""
from __future__ import annotations

from datetime import date
from io import BytesIO
import re
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
import xml.etree.ElementTree as ET

import pandas as pd

from .import_dates import initial_urgent_flags, production_completion_date


ORDER_COLUMNS = ["製令單號", "產品品號", "品名", "規格", "預計產量", "單位"]


def reveal_spec_columns(content: bytes) -> bytes:
    """Return a copy with specification columns visible; preserve other ZIP parts."""
    from openpyxl import load_workbook
    book = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = book["單頭資料"]
    targets = {cell.column for row in sheet.iter_rows(min_row=1, max_row=20) for cell in row if cell.value == "規格"}
    book.close()
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    output = BytesIO()
    with ZipFile(BytesIO(content)) as original:
        workbook = ET.fromstring(original.read("xl/workbook.xml"))
        sheet_node = next(node for node in workbook.findall("s:sheets/s:sheet", ns) if node.get("name") == "單頭資料")
        rel_id = sheet_node.get(f"{{{rel_ns}}}id")
        rels = ET.fromstring(original.read("xl/_rels/workbook.xml.rels"))
        target = next(node.get("Target") for node in rels if node.get("Id") == rel_id)
        path = target.lstrip("/") if target.startswith("/") else "xl/" + target
        original_xml = original.read(path)
        root = ET.fromstring(original_xml)
        cols = root.find("s:cols", ns)
        if cols is not None:
            for col in list(cols):
                low, high = int(col.get("min")), int(col.get("max"))
                selected = sorted(n for n in targets if low <= n <= high)
                if not selected:
                    continue
                position = list(cols).index(col)
                cols.remove(col)
                boundaries = sorted({low, high + 1, *selected, *(n + 1 for n in selected)})
                for first, end in zip(boundaries, boundaries[1:]):
                    piece = deepcopy(col)
                    piece.set("min", str(first))
                    piece.set("max", str(end - 1))
                    if first in targets:
                        piece.set("hidden", "0")
                        if float(piece.get("width", "10")) == 0:
                            piece.set("width", "20")
                    cols.insert(position, piece)
                    position += 1
        updated_xml = original_xml
        if cols is not None:
            # Replace only column definitions. Keep all original root namespace
            # declarations (including names referenced by mc:Ignorable) intact.
            updated_xml = re.sub(rb"<(?:\w+:)?cols\b[^>]*>.*?</(?:\w+:)?cols>",
                                 lambda _: ET.tostring(cols, encoding="utf-8"), original_xml, count=1, flags=re.S)
        with ZipFile(output, "w", ZIP_DEFLATED) as copied:
            for item in original.infolist():
                copied.writestr(item, updated_xml if item.filename == path else original.read(item.filename))
    return output.getvalue()


def closing_date_from_filename(filename: str) -> date:
    matches = re.findall(r"結關[\s_-]*(\d{4})[.\-/](\d{1,2})[.\-/](\d{1,2})", filename)
    if len(matches) != 1:
        raise ValueError("檔名須有唯一的結關日期，例如 結關2026.09.23；請確認後重新命名")
    try:
        return date(*map(int, matches[0]))
    except ValueError as exc:
        raise ValueError("檔名的結關日期無效") from exc


def read_exported_orders(content: bytes, filename: str) -> pd.DataFrame:
    raw = pd.read_excel(BytesIO(content), sheet_name="單頭資料", header=None, dtype=str).fillna("")
    headers = [
        index for index, row in raw.iterrows()
        if set(ORDER_COLUMNS).issubset({str(cell).strip() for cell in row})
    ]
    if len(headers) != 1:
        raise ValueError("單頭資料須有唯一標題列，包含：" + "、".join(ORDER_COLUMNS))
    header = headers[0]
    labels = raw.iloc[header].astype(str).str.strip().tolist()
    if any(labels.count(column) != 1 for column in ORDER_COLUMNS):
        raise ValueError("訂單必要欄位名稱重複，請確認來源檔案")
    rows = raw.iloc[header + 1:].copy()
    # Ignore only genuinely empty rows; incomplete orders remain in the batch.
    rows = rows.loc[rows.apply(lambda row: any(str(value).strip() for value in row), axis=1)]
    result = pd.DataFrame({column: rows.iloc[:, labels.index(column)].str.strip() for column in ORDER_COLUMNS})
    result["來源檔案"] = filename
    result["來源列"] = rows.index + 1
    result["單位"] = result["單位"].str.upper().replace({"PC": "PCS"})
    delivery = closing_date_from_filename(filename)
    result["結關日"] = pd.Timestamp(delivery)
    result["生產完成日"] = pd.Timestamp(production_completion_date(delivery))
    return result.reset_index(drop=True)


def combine_order_batches(batches: list[pd.DataFrame]) -> pd.DataFrame:
    if not batches:
        return pd.DataFrame()
    result = pd.concat(batches, ignore_index=True)
    flags = initial_urgent_flags([pd.Timestamp(value).date() for value in result["結關日"]])
    result["優先級"] = ["急單" if flag else "一般" for flag in flags]
    return result


def basic_order_issues(orders: pd.DataFrame) -> dict[int, list[str]]:
    """Row identity is preserved, including both sides of duplicate order IDs."""
    issues: dict[int, list[str]] = {}
    duplicates = orders["製令單號"].duplicated(keep=False)
    for index, row in orders.iterrows():
        reasons = []
        for column in ["製令單號", "產品品號", "品名", "規格"]:
            if not str(row[column]).strip():
                reasons.append(f"缺少{column}")
        if duplicates.loc[index] and str(row["製令單號"]).strip():
            reasons.append("製令單號重複，請確認是否重複匯入版本")
        quantity = pd.to_numeric(row["預計產量"], errors="coerce")
        if pd.isna(quantity) or quantity <= 0 or quantity == float("inf"):
            reasons.append("預計產量必須為大於零的有限數值")
        if row["單位"] != "PCS":
            reasons.append("單位僅支援 PC／PCS")
        if reasons:
            issues[index] = reasons
    return issues


def batch_must_stop(total_orders: int, problematic_orders: int) -> bool:
    if total_orders < 0 or not 0 <= problematic_orders <= total_orders:
        raise ValueError("問題工單數必須介於零與總工單數之間")
    # Integer arithmetic keeps the 30% boundary exact.
    return total_orders == 0 or problematic_orders * 10 >= total_orders * 3
