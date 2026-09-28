from io import BytesIO
from pathlib import Path

from streamlit.testing.v1 import AppTest

from aps.import_pipeline import database_signature
from test_import_pipeline import inputs


class Upload(BytesIO):
    def __init__(self, name, content):
        super().__init__(content)
        self.name = name


def test_import_ui_schedules_partial_batch_and_invalidates_replaced_orders(monkeypatch, tmp_path):
    orders, rates, molds, mappings = inputs()
    orders.loc[0, "產品品號"] = "UNKNOWN"
    cp, mp, op = Upload("capacity.xls", b"capacity"), Upload("molds.xls", b"molds"), Upload("orders.xlsx", b"orders")
    uploads = {"capacity_upload": cp, "mold_upload": mp, "orders_upload": [op]}
    monkeypatch.setattr("streamlit.file_uploader", lambda *args, **kwargs: uploads.get(kwargs.get("key")))
    monkeypatch.setattr("streamlit.delta_generator.DeltaGenerator.file_uploader", lambda *args, **kwargs: uploads.get(kwargs.get("key")))
    monkeypatch.setattr("ui.order_import.read_capacity", lambda *args: (rates, []))
    monkeypatch.setattr("ui.order_import.read_molds", lambda *args: (molds, []))
    monkeypatch.setattr("ui.order_import.read_exported_orders", lambda *args: orders.copy())
    monkeypatch.setattr("ui.order_import.reveal_spec_columns", lambda blob: blob)
    monkeypatch.setattr("aps.history.HISTORY_PATH", tmp_path / "history.json")
    app = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=30).run()
    next(r for r in app.radio if r.label == "資料來源").set_value("訂單拋轉＋資料庫").run()
    assert not app.exception
    assert next(b for b in app.button if b.label == "完成匯入，套用至排程").disabled
    app.session_state["import_mappings"] = mappings
    app.run()
    assert not app.exception
    next(b for b in app.button if b.label == "完成匯入，套用至排程").click().run()
    assert not app.exception
    next(b for b in app.button if b.label == "確認設定並開始排程").click().run()
    assert not app.exception
    assert len(app.session_state["schedule_df"]) == 10
    assert app.session_state["kpis"]["未完成 / 超出 horizon 工單數"] == 1
    assert app.session_state["import_db_signature"] == database_signature(b"capacity", b"molds")
    # A different uploaded batch must not leave the previous result executable.
    uploads["orders_upload"] = [Upload("new.xlsx", b"new")]
    app.run()
    assert not app.exception
    assert app.session_state["schedule_df"] is None
    assert next(b for b in app.button if b.label == "確認設定並開始排程").disabled


def test_unhide_copy_preserves_other_columns_and_values():
    from openpyxl import Workbook, load_workbook
    from aps.import_orders import reveal_spec_columns, read_exported_orders
    book = Workbook()
    sheet = book.active
    sheet.title = "單頭資料"
    sheet.append(["製令單號", "產品品號", "品名", "規格", "預計產量", "單位"])
    sheet.append(["WO", "P", "透明管", 'ID 1/8" x OD 1/4" x 100 FT', 1, "PC"])
    sheet.column_dimensions.group("C", "E", hidden=True)
    content = BytesIO()
    book.save(content)
    revealed = reveal_spec_columns(content.getvalue())
    copy = load_workbook(BytesIO(revealed))
    assert not copy.active.column_dimensions["D"].hidden
    assert copy.active.column_dimensions["C"].hidden
    assert copy.active.column_dimensions["E"].hidden
    assert list(sheet.values) == list(copy.active.values)
    imported = read_exported_orders(content.getvalue(), "結關2026.10.09.xlsx")
    assert imported.iloc[0]["規格"] == sheet["D2"].value
    assert imported.iloc[0]["單位"] == "PCS"
